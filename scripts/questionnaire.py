"""
Génère le questionnaire Excel à faire remplir par Emmanuelle.
Le fichier rempli s'importe dans l'administration (bouton « Importer le
questionnaire »). La colonne A de chaque onglet contient une CLÉ stable :
c'est elle que l'import lit, jamais le libellé. Ne pas renommer les onglets.
Schéma détaillé : docs/questionnaire.md

Usage : python3 scripts/questionnaire.py [sortie.xlsx]
"""
import sys, re, glob, os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation

SORTIE = sys.argv[1] if len(sys.argv) > 1 else 'docs/questionnaire-emma-broderie.xlsx'
VERSION = '1'

PRUNE = '5A2244'; ROSE = 'F4C3CF'; MENTHE = 'A9DFC5'; BEURRE = 'F5D978'; PAPIER = 'FBF7F3'; GRIS = '6E5B66'
f_titre = Font(name='Arial', size=16, bold=True, color=PRUNE)
f_entete = Font(name='Arial', size=11, bold=True, color='FFFFFF')
f_texte = Font(name='Arial', size=11)
f_aide = Font(name='Arial', size=10, italic=True, color=GRIS)
fond_entete = PatternFill('solid', fgColor=PRUNE)
fond_reponse = PatternFill('solid', fgColor='FFFBEA')
fond_cle = PatternFill('solid', fgColor='EEEEEE')
fin = Side(style='thin', color='D9CCD3')
cadre = Border(left=fin, right=fin, top=fin, bottom=fin)
haut = Alignment(vertical='top', wrap_text=True)

wb = Workbook()

def liste(ws, cellules, choix):
    dv = DataValidation(type='list', formula1='"' + ','.join(choix) + '"', allow_blank=True)
    dv.error = 'Choisissez dans la liste.'; dv.prompt = 'Choisissez dans la liste'
    ws.add_data_validation(dv); dv.add(cellules)

def entete(ws, titres, largeurs, ligne=3):
    for i, (t, l) in enumerate(zip(titres, largeurs)):
        c = ws.cell(row=ligne, column=i + 1, value=t)
        c.font = f_entete; c.fill = fond_entete; c.alignment = Alignment(vertical='center', wrap_text=True); c.border = cadre
        ws.column_dimensions[chr(65 + i)].width = l
    ws.row_dimensions[ligne].height = 30
    ws.freeze_panes = ws.cell(row=ligne + 1, column=1)

def ligne(ws, r, valeurs, reponses=()):
    for i, v in enumerate(valeurs):
        c = ws.cell(row=r, column=i + 1, value=v)
        c.font = f_texte; c.alignment = haut; c.border = cadre
        if i == 0: c.fill = fond_cle; c.font = Font(name='Arial', size=8, color='999999')
        if i in reponses: c.fill = fond_reponse

# ------------------------------------------------------------- À lire
ws = wb.active; ws.title = 'À lire'
ws.sheet_view.showGridLines = False
ws.column_dimensions['A'].width = 4; ws.column_dimensions['B'].width = 100
lignes = [
    ('Questionnaire pour votre site emma-broderie.fr', f_titre),
    ('', f_texte),
    ("Bonjour Emmanuelle,", f_texte),
    ("Ce fichier rassemble tout ce qu'il faut pour que votre site affiche vos vraies informations : vos coordonnées, vos prix, vos délais, vos avis clients.", f_texte),
    ("Une fois rempli, renvoyez-le simplement : il sera importé dans le site en un clic, sans rien retaper.", f_texte),
    ('', f_texte),
    ('Comment le remplir', Font(name='Arial', size=12, bold=True, color=PRUNE)),
    ("1. Remplissez uniquement les cases jaune pâle. Les autres colonnes donnent la question et un exemple.", f_texte),
    ("2. Quand une flèche apparaît dans la case, choisissez dans la liste (Oui / Non, etc.).", f_texte),
    ("3. Si vous ne savez pas ou ne souhaitez pas répondre, laissez la case vide : rien ne sera changé sur le site.", f_texte),
    ("4. Ne modifiez pas la première colonne grise ni le nom des onglets : ils servent à l'import.", f_texte),
    ('', f_texte),
    ('Les onglets', Font(name='Arial', size=12, bold=True, color=PRUNE)),
    ("Mes coordonnées : votre nom, votre numéro, vos délais, vos conditions.", f_texte),
    ("Mes créations : le nom et le prix « à partir de » de chaque pièce présentée sur le site.", f_texte),
    ("Mes avis : les avis de clientes que vous acceptez de publier (avec leur accord).", f_texte),
    ("Mes horaires : quand vous répondez aux messages.", f_texte),
    ("Questions : quelques questions pour que le site dise exactement ce que vous faites, ni plus ni moins.", f_texte),
    ('', f_texte),
    ("Votre adresse personnelle ne sera jamais publiée : seule la ville apparaît sur le site.", f_aide),
    (f'Version du questionnaire : {VERSION}', f_aide),
]
for i, (t, f) in enumerate(lignes, start=2):
    c = ws.cell(row=i, column=2, value=t); c.font = f; c.alignment = Alignment(wrap_text=True, vertical='top')

# ------------------------------------------------------------- Coordonnées
ws = wb.create_sheet('Mes coordonnées')
ws['B1'] = 'Mes coordonnées et mes conditions'; ws['B1'].font = f_titre
entete(ws, ['clé', 'Question', 'Votre réponse', 'Exemple ou précision'], [16, 46, 40, 50])
Q = [
    ('nom_legal', 'Votre nom et prénom tels que déclarés (mentions légales)', '', 'Emmanuelle Wagner'),
    ('siret', 'Votre numéro SIRET (14 chiffres)', '', '123 456 789 00012'),
    ('adresse_pro', "Adresse déclarée de l'entreprise (mentions légales uniquement, jamais affichée ailleurs)", '', 'Peut être une domiciliation. Obligatoire dans les mentions légales.'),
    ('ville', "Ville affichée sur le site", 'Vitré', 'Vitré'),
    ('code_postal', 'Code postal', '35500', '35500'),
    ('telephone', 'Numéro de téléphone affiché sur le site', '', '06 12 34 56 78'),
    ('whatsapp', 'Numéro WhatsApp (si différent du téléphone)', '', 'Laisser vide si identique'),
    ('email', 'Adresse e-mail pour les demandes', '', 'bonjour@emma-broderie.fr'),
    ('facebook', 'Adresse de votre page Facebook (si vous en avez une)', '', 'https://www.facebook.com/…'),
    ('annee_debut', 'En quelle année avez-vous commencé à broder ?', '', '2015'),
    ('delai_reponse', 'En combien de temps répondez-vous en général à un message ?', '', 'Dans la journée'),
    ('delai_actuel', 'Délai de réalisation actuel', '', '2 semaines'),
    ('commandes', 'Prenez-vous des commandes en ce moment ?', '', 'Ouvertes / Complet / En congés'),
    ('retour_le', 'Si en congés : date de retour', '', '15/01/2027'),
    ('date_limite_noel', 'Date limite pour commander pour Noël', '', '1er décembre'),
    ('remise_vitre', 'Remise en main propre à Vitré ?', '', 'Oui / Non'),
    ('remise_rennes', 'Remise en main propre à Rennes ?', '', 'Oui / Non'),
    ('remise_autres', "Autres communes où vous pouvez remettre une commande en main propre", '', 'Pocé-les-Bois, Châteaubourg… (laisser vide si aucune)'),
    ('envoi_france', 'Envoyez-vous par la poste partout en France ?', '', 'Oui / Non'),
    ('frais_envoi', "Frais d'envoi (montant ou règle)", '', 'Lettre suivie : 5 € ; colis : selon le poids'),
    ('emballage_cadeau', 'Proposez-vous un emballage cadeau ?', '', 'Oui / Non'),
    ('paiement', 'Moyens de paiement acceptés', '', 'Virement, espèces, chèque, Wero…'),
    ('acompte', 'Demandez-vous un acompte ? Lequel ?', '', '30 % à la validation du motif'),
    ('marches', 'Marchés ou salons où l’on peut vous rencontrer', '', 'Marché de Noël de Vitré (décembre)…'),
]
for i, (cle, q, rep, ex) in enumerate(Q, start=4):
    ligne(ws, i, [cle, q, rep or None, ex], reponses=(2,))
idx = {cle: i for i, (cle, *_ ) in enumerate(Q, start=4)}
liste(ws, f'C{idx["commandes"]}', ['Ouvertes', 'Complet', 'En congés'])
for k in ('remise_vitre', 'remise_rennes', 'envoi_france', 'emballage_cadeau'):
    liste(ws, f'C{idx[k]}', ['Oui', 'Non'])

# ------------------------------------------------------------- Créations
ws = wb.create_sheet('Mes créations')
ws['B1'] = 'Mes créations présentées sur le site'; ws['B1'].font = f_titre
ws['B2'] = "Une ligne = une carte du site. Ne touchez pas à la colonne « Carte actuelle » : modifiez seulement les cases jaunes de la même ligne."; ws['B2'].font = f_aide
entete(ws, ['clé', 'Carte actuelle (ne pas modifier)', 'Nom à afficher', 'Prix à partir de (€)', 'Délai (jours)', "Sur l'accueil ?", 'Ordre (1 = en premier)', 'Remarque'], [16, 32, 32, 14, 12, 13, 14, 40])
r = 4
def _ordre(f):
    m = re.search(r'^ordre:\s*(\d+)', open(f, encoding='utf-8').read(), re.M)
    return int(m.group(1)) if m else 100
for f in sorted(glob.glob('src/content/creations/*.md'), key=_ordre):
    t = open(f, encoding='utf-8').read()
    if re.search(r'^masquee:\s*true', t, re.M): continue
    g = lambda k: (re.search(rf'^{k}:\s*"?(.*?)"?\s*$', t, re.M) or [None, None])[1]
    slug = os.path.basename(f)[:-3]
    ligne(ws, r, [slug, g('titre'), g('titre'), int(g('prixDepart')) if g('prixDepart') else None,
                  (int(g('delaiJours')) if (g('delaiJours') or '').isdigit() else g('delaiJours')), 'Oui', r - 3, None],
          reponses=(2, 3, 4, 5, 6, 7))
    r += 1
liste(ws, f'F4:F{r - 1}', ['Oui', 'Non'])

# ------------------------------------------------------------- Avis
ws = wb.create_sheet('Mes avis')
ws['B1'] = 'Avis de clientes'; ws['B1'].font = f_titre
ws['B2'] = "Uniquement de vrais avis, avec l'accord de la personne. Le prénom seul ou l'initiale suffit."; ws['B2'].font = f_aide
entete(ws, ['clé', 'Prénom ou initiale', 'Ville', 'Date', 'Ce que la cliente a écrit', 'Création concernée', 'Accord pour publier ?'], [10, 20, 18, 14, 60, 26, 16])
for i in range(1, 11):
    ligne(ws, 3 + i, [f'avis_{i}', None, None, None, None, None, None], reponses=(1, 2, 3, 4, 5, 6))
    ws.row_dimensions[3 + i].height = 45
liste(ws, 'G4:G13', ['Oui', 'Non'])

# ------------------------------------------------------------- Horaires
ws = wb.create_sheet('Mes horaires')
ws['B1'] = 'Quand répondez-vous aux messages ?'; ws['B1'].font = f_titre
ws['B2'] = "Heures au format 9:00. Deux plages possibles pour couper le midi. Votre domicile n'est jamais présenté comme une boutique."; ws['B2'].font = f_aide
entete(ws, ['clé', 'Jour', 'Fermé ?', 'Début', 'Fin', 'Début (2e plage)', 'Fin (2e plage)'], [12, 16, 12, 12, 12, 16, 16])
for i, j in enumerate(['Lundi', 'Mardi', 'Mercredi', 'Jeudi', 'Vendredi', 'Samedi', 'Dimanche'], start=4):
    ligne(ws, i, [j.lower(), j, None, None, None, None, None], reponses=(2, 3, 4, 5, 6))
liste(ws, 'C4:C10', ['Oui', 'Non'])
c = ws.cell(row=12, column=1, value='horaires_mode'); c.fill = fond_cle; c.font = Font(name='Arial', size=8, color='999999')
ws.cell(row=12, column=2, value='Ces horaires signifient :').font = f_texte
c = ws.cell(row=12, column=3, value=None); c.fill = fond_reponse; c.border = cadre
ws.merge_cells('C12:E12')
liste(ws, 'C12', ['Je réponds aux messages', 'Retrait sur rendez-vous'])

# ------------------------------------------------------------- Questions
ws = wb.create_sheet('Questions')
ws['B1'] = 'Ce que vous faites (et ne faites pas)'; ws['B1'].font = f_titre
ws['B2'] = "Le site ne doit rien promettre que vous ne faites pas. Répondez Oui ou Non, et précisez si besoin."; ws['B2'].font = f_aide
entete(ws, ['clé', 'Question', 'Oui / Non', 'Précision'], [22, 60, 12, 50])
QQ = [
    ('dentelle_commande', 'Réalisez-vous de la dentelle libre (arbre de vie, papillon…) sur commande ?'),
    ('couture_commande', 'Cousez-vous sur commande (sacs, trousses, tabliers, pochettes) ?'),
    ('retouches', 'Faites-vous des retouches (ourlets, ajustements) ?'),
    ('vetements_mesure', 'Confectionnez-vous des vêtements sur mesure ?'),
    ('broderie_vetement_client', "Brodez-vous sur un vêtement ou un linge apporté par la cliente ?"),
    ('logos', 'Brodez-vous des logos (associations, clubs) ?'),
    ('nombre_creations', 'Environ combien de pièces avez-vous réalisées ? (un nombre)'),
    ('accord_prenoms', "Les prénoms visibles sur vos photos (Aliza, Ambre, Alicia, Lisa, Malo, Keanu, Florian, Aidan) : avez-vous l'accord des familles pour les montrer ?"),
    ('accord_portrait', "Acceptez-vous que votre photo apparaisse en haut de l'accueil ?"),
]
for i, (cle, q) in enumerate(QQ, start=4):
    ligne(ws, i, [cle, q, None, None], reponses=(2, 3))
    ws.row_dimensions[i].height = 32
liste(ws, f'C4:C{3 + len(QQ)}', ['Oui', 'Non'])

for w in wb.worksheets:
    w.sheet_properties.tabColor = {'À lire': PRUNE, 'Mes coordonnées': ROSE, 'Mes créations': BEURRE, 'Mes avis': MENTHE, 'Mes horaires': 'ADCBEB', 'Questions': GRIS}[w.title]
    w.column_dimensions['A'].hidden = w.title != 'À lire' and False
wb.properties.title = 'Questionnaire Emma Broderie'
wb.properties.keywords = f'questionnaire-emma-broderie v{VERSION}'
os.makedirs(os.path.dirname(SORTIE), exist_ok=True)
wb.save(SORTIE)
print('écrit :', SORTIE)
